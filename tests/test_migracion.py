"""Migración: propuesta, revisión en la hoja y aplicación."""

from datetime import date

import pytest
from openpyxl import load_workbook

from etl import config
from etl.modos import migracion
from etl.modos.migracion import sugerir
from tests.conftest import DATOS_REALES, TOPOGRAFO, cartera_txt

HOY = date(2026, 10, 30)


def test_sugerencias_por_pistas(tabla):
    a = {"name": "ALEJANDRIA MH SANITARIOS K0+120 14-03-2026.xlsx", "ruta": "Disco D/",
         "owners": [{"emailAddress": TOPOGRAFO}], "modifiedTime": "2026-03-20T10:00:00Z"}
    s = sugerir(a, tabla)
    assert (s["proyecto"], s["etapa"], s["frente"], s["tipo"]) == ("P001", "E01", "F01", "CMH")
    assert s["fecha"] == date(2026, 3, 14) and s["origen"] == "PEI" and s["confianza"] == "ALTA"


def test_proyecto_terminado_solo_hasta_proyecto(tabla):
    s = sugerir({"name": "mallorquin informe final.pdf", "ruta": ""}, tabla)
    assert (s["proyecto"], s["etapa"], s["frente"], s["tipo"]) == ("P009", "ZZ", "ZZ", "INF")


def test_whatsapp_fecha_del_nombre(tabla):
    s = sugerir({"name": "DOC-20260707-WA0000_", "ruta": "alejandria/"}, tabla)
    assert s["fecha"] == date(2026, 7, 7)


def _preparar(drive, tmp_path):
    mig = drive.asegurar_carpeta(config.CARPETA_MIGRACION, drive.raiz_id)
    sub = drive.asegurar_carpeta("Alejandria sanitaria", mig)
    a = drive.subir_como(cartera_txt(tmp_path / "1.txt"), sub, "POLIGONO5 14-03-2026.TXT", dueno=TOPOGRAFO)
    b = drive.subir_como(cartera_txt(tmp_path / "1.txt"), sub, "POLIGONO5 copia.TXT", dueno=TOPOGRAFO)
    return a, b


def test_propuesta_y_aplicacion(drive, db, tmp_path):
    a, b = _preparar(drive, tmp_path)
    hoja = migracion.proponer(drive, db)
    assert drive.ruta_de(hoja["id"]).startswith(config.CARPETA_REVISION_MIGRACION)

    ruta = drive.items[hoja["id"]]["_local"]
    wb = load_workbook(open(ruta, "rb"))
    ws = wb["REVISION"]
    filas = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    assert ws.cell(filas[b], 17).value == "DESCARTAR"          # copia idéntica detectada
    r = filas[a]
    assert ws.cell(r, 7).value == "P001" and ws.cell(r, 11).value == "CAR"
    ws.cell(r, 17).value = "MIGRAR"
    ws.cell(r, 19).value = "SI"
    ws.cell(filas[b], 19).value = "SI"
    wb.save(ruta)

    res = migracion.aplicar(drive, db, hoja_id=hoja["id"], hoy=HOY)
    estados = dict(zip(res.drive_id, res.resultado))
    assert estados[a] == "MIGRADO" and estados[b] == "DESCARTADO"
    oficial = db.leer_sql("SELECT documento, cod_etapa FROM reporte.v_buscador").iloc[0]
    assert oficial["documento"] == "P001-PEI-E01-F01-CAR-14-03-2026-01"
    assert db.leer_sql("SELECT via FROM ctrl.archivo WHERE estado='VIGENTE'").via[0] == "MIGRACION"
    assert db.leer_sql("SELECT * FROM reporte.v_indicadores").archivos_migrados[0] == 1
    assert any(m["name"].startswith("RESULTADO_MIGRACION_") for m in drive.items.values())


def test_sin_aprobacion_no_se_mueve(drive, db, tmp_path):
    a, _ = _preparar(drive, tmp_path)
    hoja = migracion.proponer(drive, db)
    res = migracion.aplicar(drive, db, hoja_id=hoja["id"], hoy=HOY)
    assert set(res.resultado) == {"SIN CAMBIOS"}
    assert drive.ruta_de(a).startswith(config.CARPETA_MIGRACION)


@pytest.mark.skipif(not DATOS_REALES.exists(), reason="sin carpeta de datos reales")
def test_propuesta_con_archivos_reales(drive, db):
    mig = drive.asegurar_carpeta(config.CARPETA_MIGRACION, drive.raiz_id)
    for ruta in sorted(DATOS_REALES.rglob("*")):
        if ruta.is_file():
            drive.subir_como(ruta, mig, ruta.name, dueno=TOPOGRAFO)
    hoja = migracion.proponer(drive, db)
    ws = load_workbook(open(drive.items[hoja["id"]]["_local"], "rb"))["REVISION"]
    assert ws.max_row - 1 == sum(1 for x in DATOS_REALES.rglob("*") if x.is_file())
