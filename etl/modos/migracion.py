"""
Migración de los archivos existentes (ISO 13008: planificar, revisar,
ejecutar y verificar).

PROPUESTA (manual): recorre 00_MIGRACION, sugiere para cada archivo el
nombre estándar con las pistas disponibles y genera una Hoja de cálculo
compartida en 02_REVISION_MIGRACION. Cada productor revisa sus propios
archivos (columna "responsable") y el director marca la aprobación.

APLICAR (manual): lee la hoja revisada y, solo en las filas con decisión
MIGRAR y aprobación SI, renombra y lleva el archivo a su carpeta oficial
con las mismas validaciones del modo diario. Mover y copiar se hacen en
los servidores de Google y el md5 se conserva, lo que sirve de evidencia
de que el contenido no cambió. El resultado de cada fila queda en una
hoja RESULTADO y en ctrl.migracion_item.
"""

from __future__ import annotations

import logging
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import pandas as pd

from etl import config
from etl.contenido import leer_contenido
from etl.db import BaseDatos
from etl.drive import Drive
from etl.modos.comun import Carpetas, cargar_tabla_maestra, carpeta_temporal, ejecucion
from etl.modos.diario import ProcesadorEntrada
from etl.nomenclatura import GENERAL, ORIGEN_INTERNO, TIPOS, construir_nombre
from etl.tabla_maestra import TablaMaestra, abscisa_a_metros, normalizar_texto

log = logging.getLogger("etl")

PALABRAS_TIPO = [
    ("CMH", ["mh ", "mh-", "mh_", "pozo", "pozos", "manhole", "cota tapa"]),
    ("PAC", ["preacta", "pre acta", "acta"]),
    ("PRE", ["presupuesto", "apu"]),
    ("CNT", ["cantidades", "cant ", "volumen", "volumenes", "corte y lleno", "ems cantidades"]),
    ("CAR", ["cartera", "nivelacion", "tuberia", "poligono", "polig", "coordenadas"]),
    ("MED", ["levantamiento", "medicion", "topografia"]),
    ("INF", ["informe"]),
    ("SOP", ["soporte"]),
    ("PLA", ["plano", "planta", "perfil"]),
    ("FOT", ["foto", "img_", "whatsapp image"]),
]
TIPO_POR_EXTENSION = {
    ".dwg": "PLA", ".dxf": "PLA", ".jpg": "FOT", ".jpeg": "FOT", ".png": "FOT", ".heic": "FOT",
    ".mp4": "FOT", ".mov": "FOT", ".txt": "CAR", ".csv": "CAR",
}
TIPO_POR_TABLA = {
    "control_mh": "CMH", "cartera_nivelacion": "CAR", "mediciones_topograficas": "CAR",
    "presupuesto_cantidades": "PRE", "cantidades_abscisa": "CNT", "cantidades_ejecutadas": "CNT",
}
EXTENSION_POR_MIME = {
    "application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png", "text/plain": ".txt",
    "text/csv": ".csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.ms-excel": ".xls",
}
TAMANO_MAX_PISTAS = 20 * 1024 * 1024

COLUMNAS = ["drive_id", "ruta", "nombre_actual", "tamano_mb", "propietario", "responsable",
            "proyecto", "origen", "etapa", "frente", "tipo", "fecha", "consecutivo",
            "nombre_resultante", "confianza", "pistas", "decision", "observaciones",
            "aprobado_director", "abrir_archivo"]
DECISIONES = ["MIGRAR", "DESCARTAR", "PENDIENTE"]


# =========================================================== sugerencias
def _fecha_en_texto(texto: str) -> Optional[date]:
    # 14-05-2025, 3-12-2025, 14052025 (como nombra el equipo) y 20260707 (WhatsApp)
    for patron, orden in ((r"(?<!\d)(\d{1,2})[-_.](\d{1,2})[-_.](\d{4})(?!\d)", "dma"),
                          (r"(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)", "dma"),
                          (r"(?<!\d)(20\d{2})[-_.]?(\d{2})[-_.]?(\d{2})(?!\d)", "amd")):
        for m in re.finditer(patron, texto):
            a, b, c = (int(x) for x in m.groups())
            try:
                return date(c, b, a) if orden == "dma" else date(a, b, c)
            except ValueError:
                continue
    return None


def sugerir(archivo: dict, tm: TablaMaestra, contenido=None) -> dict:
    """Propone proyecto, etapa, frente, tipo, fecha y origen con lo que se
    sabe del archivo: carpeta, nombre, extensión, dueño y, si es tabular
    y pequeño, su contenido (abscisas y coordenadas)."""
    # Guiones y guiones bajos cuentan como espacios: "VIA_LOCAL" = "via local"
    texto = re.sub(r"[_\-]+", " ", normalizar_texto(f"{archivo.get('ruta', '')} {archivo['name']}"))
    palabras = set(re.findall(r"[a-z0-9]+", texto))
    pistas, puntos = [], 0
    extension = Path(archivo["name"]).suffix.lower()

    # Coordenadas del contenido (mediana), si es una cartera
    punto = None
    if contenido is not None:
        for df, _ in contenido.tablas:
            if {"norte", "este"} <= set(df.columns):
                n = pd.to_numeric(df["norte"], errors="coerce").median()
                e = pd.to_numeric(df["este"], errors="coerce").median()
                if pd.notna(n) and pd.notna(e):
                    punto = (n, e)
                    break

    def etapas_con_punto(lista):
        if punto is None:
            return []
        return [e for e in lista if None not in (e.norte_min, e.norte_max, e.este_min, e.este_max)
                and e.norte_min <= punto[0] <= e.norte_max and e.este_min <= punto[1] <= e.este_max]

    # Contratistas mencionados en la carpeta o el nombre (código o nombre)
    contratistas = {c for c, datos in tm.contratistas.items()
                    if c.lower() in palabras or any(
                        w in palabras for w in re.findall(r"[a-z0-9]{2,}", normalizar_texto(datos["nombre"])))}

    # Proyecto: palabra clave; si no hay, por coordenadas
    proyecto = None
    candidatos = []
    for p in tm.proyectos.values():
        claves = p.palabras_clave + [normalizar_texto(p.nombre)]
        if p.codigo.lower() in palabras or any(k and k in texto for k in claves):
            candidatos.append(p.codigo)
    if len(candidatos) == 1:
        proyecto = candidatos[0]
        pistas.append(f"proyecto por palabra clave ({proyecto})")
        puntos += 1
    elif len(tm.proyectos) == 1:
        proyecto = next(iter(tm.proyectos))
    elif len(candidatos) > 1:
        pistas.append(f"varios proyectos posibles: {', '.join(candidatos)}")
    if proyecto is None and not candidatos:
        por_coord = sorted({e.proyecto for e in etapas_con_punto(tm.etapas.values())})
        if len(por_coord) == 1:
            proyecto = por_coord[0]
            pistas.append(f"proyecto por coordenadas ({proyecto})")
            puntos += 1

    # Etapa: abscisa, coordenadas, palabra clave o contratista
    etapa = GENERAL
    terminado = proyecto and tm.proyectos[proyecto].estado == "TERMINADO"
    if proyecto and not terminado:
        etapas = [e for (p, _), e in tm.etapas.items() if p == proyecto]
        por_clave = [e.etapa for e in etapas if any(k and k in texto for k in e.palabras_clave)]
        abscisas = [abscisa_a_metros(m) for m in re.findall(r"k\s*\d+\s*\+\s*\d+(?:[.,]\d+)?", texto)]
        por_abscisa = [e.etapa for e in etapas for a in abscisas
                       if a is not None and e.abscisa_inicio is not None and e.abscisa_fin is not None
                       and e.abscisa_inicio <= a <= e.abscisa_fin]
        por_coord = [e.etapa for e in etapas_con_punto(etapas)]
        por_contratista = [a.etapa for a in tm.frentes if a.proyecto == proyecto and a.contratista in contratistas]
        for fuente, lista in (("abscisa", por_abscisa), ("coordenadas", por_coord),
                              ("palabra clave", por_clave), ("contratista", por_contratista)):
            unicos = sorted(set(lista))
            if len(unicos) == 1:
                etapa = unicos[0]
                pistas.append(f"etapa por {fuente} ({etapa})")
                puntos += 1
                break
    elif terminado:
        pistas.append("proyecto terminado: se clasifica solo hasta proyecto")

    # Frente: el único de la etapa, o el del contratista mencionado en la ruta
    # (es común que el equipo tenga una carpeta por contratista)
    frente = GENERAL
    if etapa != GENERAL:
        asignaciones = [a for a in tm.frentes if (a.proyecto, a.etapa) == (proyecto, etapa)]
        frentes = sorted({a.frente for a in asignaciones})
        if len(frentes) == 1:
            frente = frentes[0]
        else:
            por_contratista = sorted({a.frente for a in asignaciones if a.contratista in contratistas})
            if len(por_contratista) == 1:
                frente = por_contratista[0]
                pistas.append(f"frente por contratista en la ruta ({frente})")

    # Tipo
    tipo = None
    if contenido is not None:
        tablas = {TIPO_POR_TABLA[t] for _, t in contenido.tablas if t in TIPO_POR_TABLA}
        if len(tablas) == 1:
            tipo = tablas.pop()
            pistas.append(f"tipo por contenido ({tipo})")
            puntos += 1
    if not tipo:
        for codigo, palabras in PALABRAS_TIPO:
            if any(p in f" {texto} " for p in palabras):
                tipo = codigo
                pistas.append(f"tipo por nombre ({tipo})")
                puntos += 1
                break
    if not tipo and extension in TIPO_POR_EXTENSION:
        tipo = TIPO_POR_EXTENSION[extension]
        pistas.append(f"tipo por extensión ({tipo})")
    if not tipo and extension == ".pdf":
        pistas.append("PDF sin pistas de tipo: elegir INF, PLA, SOP o PAC")

    # Fecha
    fecha = _fecha_en_texto(archivo["name"]) or _fecha_en_texto(archivo.get("ruta", ""))
    if fecha:
        pistas.append("fecha tomada del nombre")
    else:
        modif = archivo.get("modifiedTime") or archivo.get("createdTime")
        fecha = datetime.fromisoformat(modif.replace("Z", "+00:00")).date() if modif else date.today()
        pistas.append("fecha de modificación en Drive, confirmar")

    # Origen: por el correo del dueño
    dueno = ((archivo.get("owners") or [{}])[0].get("emailAddress") or "").lower()
    origen = tm.origen_de_correo(dueno) or ORIGEN_INTERNO

    confianza = "ALTA" if (proyecto and puntos >= 3) else "MEDIA" if (proyecto and tipo) else "BAJA"
    return {"proyecto": proyecto or "", "origen": origen, "etapa": etapa, "frente": frente,
            "tipo": tipo or "", "fecha": fecha, "confianza": confianza, "pistas": "; ".join(pistas),
            "propietario": dueno}


def _responsable(archivo: dict, sugerencia: dict, cuenta_empresa: str) -> str:
    """Quién revisa la fila: el dueño del archivo; si lo subió la cuenta de la
    empresa, el nombre de la subcarpeta de 00_MIGRACION (ej. 'Andres/')."""
    dueno = sugerencia.get("propietario") or ""
    if dueno and dueno != cuenta_empresa:
        return dueno
    primera = (archivo.get("ruta") or "").split("/")[0]
    return primera or "director"


def _nombre_parecido(nombre: str) -> str:
    """'ALEJANDRIA (1).CSV' y 'alejandria copia.csv' -> 'alejandria'"""
    base = normalizar_texto(Path(nombre).stem)
    base = re.sub(r"\(\d+\)|\bcopia( de)?\b|\bcopy\b|\bfinal\b|\bv\d+\b", " ", base)
    return re.sub(r"[\s_\-.]+", " ", base).strip()


# ============================================================== hoja Excel
def construir_hoja(filas: list, tm: TablaMaestra, ruta: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ins = wb.active
    ins.title = "INSTRUCCIONES"
    for i, l in enumerate([
        "REVISIÓN DE LA MIGRACIÓN. La ETL propuso un nombre estándar para cada archivo de 00_MIGRACION.",
        "1. Filtre la columna 'responsable' por su correo: cada persona revisa los archivos que subió.",
        "2. Corrija proyecto, etapa, frente, tipo o fecha si la sugerencia no es correcta (use las listas).",
        "3. En 'decision' elija MIGRAR, DESCARTAR (duplicados o basura) o PENDIENTE.",
        "4. El director revisa y escribe SI en 'aprobado_director'. Solo se migran filas con MIGRAR y SI.",
        "5. 'consecutivo' vacío = la ETL asigna el siguiente libre de ese día.",
        "6. La fecha va como DD-MM-AAAA. Confianza BAJA = revisar con más cuidado.",
        "No borre ni cambie las columnas drive_id, ruta ni nombre_actual.",
    ], 1):
        ins.cell(row=i, column=1, value=l)
    ins.column_dimensions["A"].width = 110

    listas = wb.create_sheet("LISTAS")
    etapas = sorted({e for (_, e) in tm.etapas} | {GENERAL})
    frentes = sorted({a.frente for a in tm.frentes} | {GENERAL})
    origenes = [ORIGEN_INTERNO] + sorted(tm.contratistas)
    columnas_listas = {"A": sorted(tm.proyectos), "B": origenes, "C": etapas, "D": frentes,
                       "E": list(TIPOS), "F": DECISIONES, "G": ["SI", "NO"]}
    for col, valores in columnas_listas.items():
        for i, v in enumerate(valores, 1):
            listas[f"{col}{i}"] = v

    ws = wb.create_sheet("REVISION", 0)
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True)
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for i, f in enumerate(filas, 2):
        ws.append([f["drive_id"], f["ruta"], f["nombre_actual"], f["tamano_mb"], f["propietario"],
                   f["responsable"], f["proyecto"], f["origen"], f["etapa"], f["frente"], f["tipo"],
                   f["fecha"].strftime("%d-%m-%Y"), None,
                   f'=G{i}&"-"&H{i}&"-"&I{i}&"-"&J{i}&"-"&K{i}&"-"&L{i}&"-"&IF(M{i}="","NN",TEXT(M{i},"00"))',
                   f["confianza"], f["pistas"], f["decision"], f["observaciones"], "",
                   f'=HYPERLINK("{f["enlace"]}","abrir")' if f.get("enlace") else ""])
        ws[f"L{i}"].number_format = "@"
        for col in "GHIJKLMQS":
            ws[f"{col}{i}"].fill = amarillo
    n = max(len(filas) + 1, 2)
    for col_hoja, col_lista, cuantos in (("G", "A", len(tm.proyectos)), ("H", "B", len(origenes)),
                                         ("I", "C", len(etapas)), ("J", "D", len(frentes)),
                                         ("K", "E", len(TIPOS)), ("Q", "F", len(DECISIONES)),
                                         ("S", "G", 2)):
        dv = DataValidation(type="list", formula1=f"=LISTAS!${col_lista}$1:${col_lista}${max(cuantos, 1)}",
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_hoja}2:{col_hoja}{n}")
    anchos = {"A": 12, "B": 30, "C": 40, "D": 9, "E": 26, "F": 26, "N": 40, "P": 45, "R": 30}
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:T{n}"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


# ================================================================ propuesta
def proponer(drive: Optional[Drive] = None, db: Optional[BaseDatos] = None):
    drive = drive or Drive()
    db = db or BaseDatos()
    with ejecucion(db, "migracion-propuesta") as (ej_id, c), carpeta_temporal() as tmp:
        tmp = Path(tmp)
        raiz = drive.raiz()
        tm = cargar_tabla_maestra(drive, raiz, tmp, db, ej_id)
        db.sincronizar_tabla_maestra(tm)
        carpetas = Carpetas(drive, raiz, tm)
        origen_id = carpetas.control(config.CARPETA_MIGRACION)
        archivos = drive.listar_archivos(origen_id, excluir={"_DESCARTADOS"})
        log.info("00_MIGRACION: %d archivo(s)", len(archivos))

        yo = drive.yo()
        vistos_md5, vistos_nombre = {}, {}
        filas = []
        for a in sorted(archivos, key=lambda x: (x.get("ruta", ""), x["name"])):
            c["revisados"] += 1
            contenido = None
            ext = Path(a["name"]).suffix.lower()
            tam = int(a.get("size") or 0)
            if (ext in config.EXTENSIONES_A_TRANSFORMAR or not ext) and 0 < tam <= TAMANO_MAX_PISTAS:
                try:
                    ruta = drive.descargar(a, tmp / f"m_{a['id']}{ext}")
                    contenido = leer_contenido(ruta, a["name"])
                    ruta.unlink(missing_ok=True)
                except Exception as exc:  # noqa: BLE001
                    log.debug("  sin pistas de contenido para %s: %s", a["name"], exc)
            s = sugerir(a, tm, contenido)
            decision, obs = "PENDIENTE", ""
            md5 = a.get("md5Checksum")
            if md5 and md5 in vistos_md5:
                decision, obs = "DESCARTAR", f"contenido idéntico a {vistos_md5[md5]}"
                c["duplicados"] += 1
            elif db.oficial_por_md5(md5):
                decision, obs = "DESCARTAR", "ya existe en lo oficial con el mismo contenido"
                c["duplicados"] += 1
            elif _nombre_parecido(a["name"]) in vistos_nombre:
                obs = (f"posible copia o versión de {vistos_nombre[_nombre_parecido(a['name'])]}; "
                       "si es una corrección, migre solo la última o use el mismo consecutivo")
            elif s["confianza"] == "ALTA":
                decision = "MIGRAR"
            vistos_nombre.setdefault(_nombre_parecido(a["name"]), f"{a.get('ruta', '')}{a['name']}")
            if md5:
                vistos_md5.setdefault(md5, f"{a.get('ruta', '')}{a['name']}")
            fila = {"drive_id": a["id"], "ruta": a.get("ruta", ""), "nombre_actual": a["name"],
                    "tamano_mb": round(tam / 1048576, 2), "responsable": _responsable(a, s, yo),
                    "enlace": a.get("webViewLink", ""),
                    "decision": decision, "observaciones": obs, **s}
            filas.append(fila)
            db.registrar_migracion_item(ej_id, {
                "d": a["id"], "r": fila["ruta"], "n": a["name"], "m": md5, "t": tam or None,
                "p": s["propietario"], "s": construir_nombre(s["proyecto"] or "P???", s["origen"], s["etapa"],
                                                             s["frente"], s["tipo"] or "???", s["fecha"], 1),
                "c": s["confianza"], "pi": s["pistas"]})

        nombre_hoja = f"REVISION_MIGRACION_{date.today():%Y-%m-%d}"
        local = construir_hoja(filas, tm, tmp / f"{nombre_hoja}.xlsx")
        destino = carpetas.control(config.CARPETA_REVISION_MIGRACION)
        hoja = drive.subir(local, destino, nombre_hoja, como_hoja=True)
        db.evento(ej_id, "migracion_propuesta", hoja["id"], nombre_hoja,
                  f"{len(filas)} filas; {Counter(f['confianza'] for f in filas)}")
        log.debug("Hoja de revisión: %s", hoja.get("webViewLink"))
        c["por_revisar"] = len(filas)
        return hoja


# ================================================================== aplicar
def _texto(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip().upper()


def _extension(nombre: str, mime: str) -> str:
    ext = Path(nombre).suffix.lower()
    return ext if ext else EXTENSION_POR_MIME.get(mime, "")


def aplicar(drive: Optional[Drive] = None, db: Optional[BaseDatos] = None, hoja_id: Optional[str] = None,
            hoy: Optional[date] = None):
    drive = drive or Drive()
    db = db or BaseDatos()
    with ejecucion(db, "migracion-aplicar") as (ej_id, c), carpeta_temporal() as tmp:
        tmp = Path(tmp)
        raiz = drive.raiz()
        tm = cargar_tabla_maestra(drive, raiz, tmp, db, ej_id)
        db.sincronizar_tabla_maestra(tm)
        proc = ProcesadorEntrada(drive, db, tm, raiz, ej_id, c, tmp, hoy, mover_apartados=False)

        if hoja_id:
            hoja = drive.obtener(hoja_id)
        else:
            carpeta = proc.carpetas.control(config.CARPETA_REVISION_MIGRACION)
            hojas = [h for h in drive.listar_archivos(carpeta, recursivo=False)
                     if h["name"].startswith("REVISION_MIGRACION_")]
            if not hojas:
                raise RuntimeError("No hay hojas REVISION_MIGRACION_* en 02_REVISION_MIGRACION.")
            hoja = max(hojas, key=lambda h: h["name"])
        log.info("Aplicando la hoja %s", hoja["name"])
        local = drive.exportar_xlsx(hoja, tmp / "revision.xlsx")
        df = pd.read_excel(local, sheet_name="REVISION", dtype=object)

        consecutivos = defaultdict(int)
        resultados = []
        for _, f in df.iterrows():
            drive_id = _texto(f.get("drive_id")) and str(f.get("drive_id")).strip()
            decision, aprobado = _texto(f.get("decision")), _texto(f.get("aprobado_director"))
            if not drive_id:
                continue
            if decision == "DESCARTAR" and aprobado == "SI":
                db.actualizar_migracion_item(drive_id, "DESCARTADO", aprobado_por="director")
                resultados.append((drive_id, "DESCARTADO", ""))
                continue
            if decision != "MIGRAR" or aprobado != "SI":
                resultados.append((drive_id, "SIN CAMBIOS", "falta MIGRAR + SI"))
                continue
            try:
                archivo = drive.obtener(drive_id)
                partes = {k: _texto(f.get(k)) for k in ("proyecto", "origen", "etapa", "frente", "tipo", "fecha")}
                fecha = datetime.strptime(partes["fecha"], "%d-%m-%Y").date()
                ext = _extension(archivo["name"], archivo["mimeType"])
                clave = (partes["proyecto"], partes["origen"], partes["etapa"], partes["frente"], partes["tipo"], fecha)
                consecutivo = _texto(f.get("consecutivo"))
                if consecutivo:
                    numero = int(float(consecutivo))
                else:
                    if clave not in consecutivos:
                        consecutivos[clave] = db.siguiente_consecutivo(*clave) - 1
                    consecutivos[clave] += 1
                    numero = consecutivos[clave]
                nombre = construir_nombre(*clave, numero, ext)
                archivo["ruta"] = _texto(f.get("ruta"))
                estado = proc.procesar(archivo, via="MIGRACION", nombre_forzado=nombre)
                if estado == "VIGENTE":
                    db.actualizar_migracion_item(drive_id, "MIGRADO", nombre, "director")
                    resultados.append((drive_id, "MIGRADO", nombre))
                else:
                    db.actualizar_migracion_item(drive_id, "ERROR", nombre)
                    resultados.append((drive_id, estado, proc.ultimo_motivo))
            except Exception as exc:  # noqa: BLE001
                c["errores"] += 1
                db.actualizar_migracion_item(drive_id, "ERROR")
                db.evento(ej_id, "error", drive_id, None, f"migración: {exc}")
                resultados.append((drive_id, "ERROR", str(exc)))

        res = pd.DataFrame(resultados, columns=["drive_id", "resultado", "detalle"])
        salida = df.merge(res, on="drive_id", how="left")
        nombre_res = hoja["name"].replace("REVISION_", "RESULTADO_") + f"_{datetime.now():%H%M}"
        ruta_res = tmp / f"{nombre_res}.xlsx"
        salida.to_excel(ruta_res, index=False, sheet_name="RESULTADO")
        drive.subir(ruta_res, hoja["parents"][0] if hoja.get("parents") else raiz, nombre_res, como_hoja=True)
        log.info("Resultado: %s", res["resultado"].value_counts().to_dict())
        return res
