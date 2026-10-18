"""
Tabla maestra: la hoja de Drive que mantiene el director y que es el
centro del gobierno de datos. La ETL la lee en cada corrida (exportada
como .xlsx, sin descargar nada más) y con ella valida cada archivo.

Hojas y columnas esperadas (la plantilla se genera con
`python -m etl plantilla`):

  PROYECTOS     codigo, nombre, estado, fecha_inicio, fecha_fin, palabras_clave
  ETAPAS        proyecto, etapa, nombre, fecha_inicio, fecha_fin,
                abscisa_inicio, abscisa_fin, norte_min, norte_max,
                este_min, este_max, palabras_clave
  FRENTES       proyecto, etapa, frente, nombre, contratista, fecha_inicio, fecha_fin
  CONTRATISTAS  codigo, nombre, correos
  PERSONAS      correo, nombre, rol, origen

Regla de FRENTES: si un frente cambia de contratista se agrega una
fila nueva con sus fechas, sin borrar la anterior. Así la ETL sabe qué
contratista estaba a cargo en la fecha de cada archivo.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd

from etl.nomenclatura import GENERAL, ORIGEN_INTERNO, NombreValidado, TIPOS

HOJAS = {
    "PROYECTOS": ["codigo", "nombre", "estado", "fecha_inicio", "fecha_fin", "palabras_clave"],
    "ETAPAS": ["proyecto", "etapa", "nombre", "fecha_inicio", "fecha_fin",
               "abscisa_inicio", "abscisa_fin", "norte_min", "norte_max",
               "este_min", "este_max", "palabras_clave"],
    "FRENTES": ["proyecto", "etapa", "frente", "nombre", "contratista",
                "fecha_inicio", "fecha_fin"],
    "CONTRATISTAS": ["codigo", "nombre", "correos"],
    "PERSONAS": ["correo", "nombre", "rol", "origen"],
}


def normalizar_texto(texto) -> str:
    """Minúsculas y sin tildes, para comparar palabras clave."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in t if not unicodedata.combining(c)).lower().strip()


def abscisa_a_metros(valor) -> Optional[float]:
    """'K0+470' -> 470.0 ; 'k1+020.5' -> 1020.5 ; 470 -> 470.0"""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)) or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    m = re.match(r"^\s*[kK]\s*(\d+)\s*\+\s*(\d+(?:[.,]\d+)?)\s*$", str(valor))
    if m:
        return int(m.group(1)) * 1000 + float(m.group(2).replace(",", "."))
    try:
        return float(str(valor).replace(",", "."))
    except ValueError:
        return None


def _fecha(valor) -> Optional[date]:
    if valor is None or valor == "" or valor == "-" or (isinstance(valor, float) and pd.isna(valor)):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, pd.Timestamp):
        return valor.date()
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    return None


def _num(valor) -> Optional[float]:
    try:
        f = float(valor)
        return None if pd.isna(f) else f
    except (TypeError, ValueError):
        return None


def _txt(valor) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    return str(valor).strip()


@dataclass
class Proyecto:
    codigo: str
    nombre: str
    estado: str = "ACTIVO"
    fecha_inicio: Optional[date] = None
    fecha_fin: Optional[date] = None
    palabras_clave: List[str] = field(default_factory=list)


@dataclass
class Etapa:
    proyecto: str
    etapa: str
    nombre: str
    fecha_inicio: Optional[date] = None
    fecha_fin: Optional[date] = None
    abscisa_inicio: Optional[float] = None
    abscisa_fin: Optional[float] = None
    norte_min: Optional[float] = None
    norte_max: Optional[float] = None
    este_min: Optional[float] = None
    este_max: Optional[float] = None
    palabras_clave: List[str] = field(default_factory=list)


@dataclass
class AsignacionFrente:
    proyecto: str
    etapa: str
    frente: str
    nombre: str
    contratista: str
    fecha_inicio: Optional[date] = None
    fecha_fin: Optional[date] = None

    def vigente_en(self, dia: date) -> bool:
        if self.fecha_inicio and dia < self.fecha_inicio:
            return False
        if self.fecha_fin and dia > self.fecha_fin:
            return False
        return True


@dataclass
class TablaMaestra:
    proyectos: Dict[str, Proyecto] = field(default_factory=dict)
    etapas: Dict[Tuple[str, str], Etapa] = field(default_factory=dict)
    frentes: List[AsignacionFrente] = field(default_factory=list)
    contratistas: Dict[str, dict] = field(default_factory=dict)
    personas: Dict[str, dict] = field(default_factory=dict)  # correo -> datos
    problemas: List[str] = field(default_factory=list)       # errores de la propia tabla

    # ---------------------------------------------------------------- carga
    @classmethod
    def desde_excel(cls, ruta: Path) -> "TablaMaestra":
        hojas = pd.read_excel(ruta, sheet_name=None, dtype=object)
        hojas = {k.strip().upper(): v for k, v in hojas.items()}
        t = cls()
        for nombre, columnas in HOJAS.items():
            if nombre not in hojas:
                t.problemas.append(f"falta la hoja {nombre}")
                hojas[nombre] = pd.DataFrame(columns=columnas)
            df = hojas[nombre]
            df.columns = [normalizar_texto(c).replace(" ", "_") for c in df.columns]
            faltan = [c for c in columnas if c not in df.columns]
            if faltan and nombre in ("PROYECTOS", "ETAPAS", "FRENTES"):
                t.problemas.append(f"a la hoja {nombre} le faltan las columnas {', '.join(faltan)}")
            for c in faltan:
                df[c] = None
            hojas[nombre] = df.dropna(how="all")

        def claves(v):
            return [normalizar_texto(p) for p in _txt(v).split(",") if p.strip()]

        for _, f in hojas["PROYECTOS"].iterrows():
            codigo = _txt(f["codigo"]).upper()
            if not codigo:
                continue
            t.proyectos[codigo] = Proyecto(
                codigo=codigo, nombre=_txt(f["nombre"]),
                estado=_txt(f["estado"]).upper() or "ACTIVO",
                fecha_inicio=_fecha(f["fecha_inicio"]), fecha_fin=_fecha(f["fecha_fin"]),
                palabras_clave=claves(f["palabras_clave"]),
            )

        for _, f in hojas["ETAPAS"].iterrows():
            p, e = _txt(f["proyecto"]).upper(), _txt(f["etapa"]).upper()
            if not p or not e:
                continue
            if p not in t.proyectos:
                t.problemas.append(f"la etapa {p}-{e} es de un proyecto que no está en PROYECTOS")
            t.etapas[(p, e)] = Etapa(
                proyecto=p, etapa=e, nombre=_txt(f["nombre"]),
                fecha_inicio=_fecha(f["fecha_inicio"]), fecha_fin=_fecha(f["fecha_fin"]),
                abscisa_inicio=abscisa_a_metros(f["abscisa_inicio"]),
                abscisa_fin=abscisa_a_metros(f["abscisa_fin"]),
                norte_min=_num(f["norte_min"]), norte_max=_num(f["norte_max"]),
                este_min=_num(f["este_min"]), este_max=_num(f["este_max"]),
                palabras_clave=claves(f["palabras_clave"]),
            )

        for _, f in hojas["CONTRATISTAS"].iterrows():
            codigo = _txt(f["codigo"]).upper()
            if codigo:
                t.contratistas[codigo] = {
                    "nombre": _txt(f["nombre"]),
                    "correos": [c.strip().lower() for c in _txt(f["correos"]).split(",") if c.strip()],
                }

        for _, f in hojas["FRENTES"].iterrows():
            p, e, fr = (_txt(f[c]).upper() for c in ("proyecto", "etapa", "frente"))
            if not (p and e and fr):
                continue
            contratista = _txt(f["contratista"]).upper()
            if (p, e) not in t.etapas:
                t.problemas.append(f"el frente {p}-{e}-{fr} es de una etapa que no está en ETAPAS")
            if contratista and contratista not in t.contratistas:
                t.problemas.append(f"el frente {p}-{e}-{fr} tiene un contratista '{contratista}' que no está en CONTRATISTAS")
            t.frentes.append(AsignacionFrente(
                proyecto=p, etapa=e, frente=fr, nombre=_txt(f["nombre"]),
                contratista=contratista,
                fecha_inicio=_fecha(f["fecha_inicio"]), fecha_fin=_fecha(f["fecha_fin"]),
            ))

        for _, f in hojas["PERSONAS"].iterrows():
            correo = _txt(f["correo"]).lower()
            if correo:
                t.personas[correo] = {"nombre": _txt(f["nombre"]), "rol": _txt(f["rol"]),
                                      "origen": _txt(f["origen"]).upper() or ORIGEN_INTERNO}
        # Los correos de los contratistas también identifican el origen
        for codigo, datos in t.contratistas.items():
            for correo in datos["correos"]:
                t.personas.setdefault(correo, {"nombre": datos["nombre"], "rol": "contratista",
                                               "origen": codigo})
        return t

    # --------------------------------------------------------- consultas
    def contratista_en(self, proyecto, etapa, frente, dia: date) -> Optional[str]:
        for a in self.frentes:
            if (a.proyecto, a.etapa, a.frente) == (proyecto, etapa, frente) and a.vigente_en(dia):
                return a.contratista
        return None

    def contratistas_activos(self, proyecto, etapa: Optional[str], dia: date) -> set:
        return {a.contratista for a in self.frentes
                if a.proyecto == proyecto and (etapa is None or a.etapa == etapa)
                and a.vigente_en(dia) and a.contratista}

    def origen_de_correo(self, correo: Optional[str]) -> Optional[str]:
        if not correo:
            return None
        datos = self.personas.get(correo.lower())
        return datos["origen"] if datos else None

    def nombre_proyecto(self, codigo):
        p = self.proyectos.get(codigo)
        return p.nombre if p else ""

    def nombre_etapa(self, proyecto, etapa):
        if etapa == GENERAL:
            return "General"
        e = self.etapas.get((proyecto, etapa))
        return e.nombre if e else ""

    # --------------------------------------------------------- validación
    def validar(self, n: NombreValidado) -> Tuple[List[str], List[str], Optional[str]]:
        """Revisa un nombre ya validado en forma contra la tabla maestra.
        Devuelve (motivos_de_rechazo, avisos, contratista_responsable)."""
        motivos, avisos = [], []
        p, e, fr, o, dia = n.proyecto, n.etapa, n.frente, n.origen, n.fecha

        if p not in self.proyectos:
            return [f"el proyecto {p} no está registrado en la tabla maestra"], avisos, None
        if e != GENERAL and (p, e) not in self.etapas:
            return [f"la etapa {p}-{e} no está registrada en la tabla maestra"], avisos, None
        if e == GENERAL and fr != GENERAL:
            return ["si la etapa es ZZ el frente también debe ser ZZ"], avisos, None
        if fr != GENERAL and not any((a.proyecto, a.etapa, a.frente) == (p, e, fr) for a in self.frentes):
            return [f"frente no registrado: {p}-{e}-{fr}"], avisos, None

        contratista = None
        if fr != GENERAL:
            contratista = self.contratista_en(p, e, fr, dia)
            if contratista is None:
                avisos.append(f"el frente {p}-{e}-{fr} no tenía contratista asignado el {dia:%d-%m-%Y}")

        if o != ORIGEN_INTERNO:
            if o not in self.contratistas:
                motivos.append(f"el origen {o} no es PEI ni un contratista registrado")
            elif fr != GENERAL:
                if contratista and o != contratista:
                    motivos.append(
                        f"el origen {o} no corresponde al contratista a cargo del frente "
                        f"{p}-{e}-{fr} el {dia:%d-%m-%Y}, que era {contratista}")
            else:
                activos = self.contratistas_activos(p, None if e == GENERAL else e, dia)
                if o not in activos:
                    motivos.append(f"el contratista {o} no tenía frentes en {p}"
                                   f"{'' if e == GENERAL else '-' + e} el {dia:%d-%m-%Y}")

        if e != GENERAL:
            et = self.etapas[(p, e)]
            if (et.fecha_inicio and dia < et.fecha_inicio) or (et.fecha_fin and dia > et.fecha_fin):
                avisos.append(f"la fecha {dia:%d-%m-%Y} está fuera de las fechas de la etapa {p}-{e}")
        return motivos, avisos, contratista or (o if o != ORIGEN_INTERNO else None)


# Qué va en cada columna: se usa para la hoja INSTRUCCIONES y para las notas
# de los encabezados. (hoja, columna, qué escribir, ejemplo, obligatoria)
GUIA_COLUMNAS = [
    ("CONTRATISTAS", "codigo", "Tres letras que identifican al contratista. Es el código que usa en el nombre de los archivos que entrega. No se cambia nunca.", "EMS", "Sí"),
    ("CONTRATISTAS", "nombre", "Nombre del contratista.", "EMS Ingeniería S.A.S.", "Sí"),
    ("CONTRATISTAS", "correos", "Correos de Google con los que el contratista sube archivos, separados por coma. Sirven para reconocer quién subió cada archivo.", "obra@ems.com", "No"),
    ("PROYECTOS", "codigo", "P y tres números. Cada proyecto tiene uno propio que no se repite ni se cambia.", "P001", "Sí"),
    ("PROYECTOS", "nombre", "Nombre del proyecto como lo conoce el equipo.", "Alejandría", "Sí"),
    ("PROYECTOS", "estado", "ACTIVO, TERMINADO o SUSPENDIDO (elegir de la lista). Los proyectos terminados solo se clasifican hasta proyecto.", "ACTIVO", "Sí"),
    ("PROYECTOS", "fecha_inicio", "Fecha de inicio, en formato AAAA-MM-DD.", "2025-05-01", "Sí"),
    ("PROYECTOS", "fecha_fin", "Fecha de terminación. Vacía si el proyecto sigue en curso.", "2026-12-31", "No"),
    ("PROYECTOS", "palabras_clave", "Palabras con las que el equipo nombra el proyecto en carpetas y archivos, separadas por coma. Ayudan a clasificar los archivos antiguos.", "alejandria, calle 100", "Recomendada"),
    ("ETAPAS", "proyecto", "Código del proyecto al que pertenece la etapa. Debe existir en PROYECTOS.", "P001", "Sí"),
    ("ETAPAS", "etapa", "E y dos números. Se puede repetir entre proyectos distintos.", "E04", "Sí"),
    ("ETAPAS", "nombre", "Nombre de la etapa.", "Etapa 4 - Sector 5", "Sí"),
    ("ETAPAS", "fecha_inicio", "Fecha de inicio de la etapa, AAAA-MM-DD.", "2025-05-01", "Sí"),
    ("ETAPAS", "fecha_fin", "Fecha de fin. Vacía si sigue en curso.", "", "No"),
    ("ETAPAS", "abscisa_inicio", "Abscisa donde empieza la etapa.", "K0+000", "No"),
    ("ETAPAS", "abscisa_fin", "Abscisa donde termina la etapa.", "K0+800", "No"),
    ("ETAPAS", "norte_min", "Coordenada norte mínima de la zona de la etapa. Con las cuatro coordenadas el sistema ubica las carteras antiguas aunque su nombre no diga nada.", "1711000", "No"),
    ("ETAPAS", "norte_max", "Coordenada norte máxima de la zona.", "1711550", "No"),
    ("ETAPAS", "este_min", "Coordenada este mínima de la zona.", "918500", "No"),
    ("ETAPAS", "este_max", "Coordenada este máxima de la zona.", "919250", "No"),
    ("ETAPAS", "palabras_clave", "Palabras con las que el equipo nombra la etapa, separadas por coma.", "etapa 4, sector 5", "Recomendada"),
    ("FRENTES", "proyecto", "Código del proyecto. Debe existir en PROYECTOS.", "P001", "Sí"),
    ("FRENTES", "etapa", "Código de la etapa. Debe existir en ETAPAS para ese proyecto.", "E04", "Sí"),
    ("FRENTES", "frente", "F y dos números.", "F01", "Sí"),
    ("FRENTES", "nombre", "Qué se hace en ese frente.", "Redes sanitarias", "Sí"),
    ("FRENTES", "contratista", "Código del contratista a cargo, tomado de CONTRATISTAS. Vacío si lo ejecuta PEI.", "EMS", "No"),
    ("FRENTES", "fecha_inicio", "Desde cuándo está ese contratista a cargo, AAAA-MM-DD.", "2025-05-01", "Sí"),
    ("FRENTES", "fecha_fin", "Hasta cuándo. Vacía si sigue a cargo. Si el frente cambia de contratista, se pone la fecha fin aquí y se agrega una fila nueva.", "", "No"),
    ("PERSONAS", "correo", "Cuenta de Google con la que la persona sube archivos al Drive.", "topografo@gmail.com", "Sí"),
    ("PERSONAS", "nombre", "Nombre de la persona.", "Andrés Pérez", "Sí"),
    ("PERSONAS", "rol", "director, suplente, residente, topografo o auxiliar. El director y el suplente pueden reemplazar archivos oficiales de otras personas.", "topografo", "Sí"),
    ("PERSONAS", "origen", "PEI si trabaja para la interventoría, o el código del contratista si trabaja para uno.", "PEI", "Sí"),
]

PROPOSITO_HOJAS = {
    "CONTRATISTAS": "Las empresas que ejecutan la obra. Se llena primero porque FRENTES usa sus códigos.",
    "PROYECTOS": "Cada obra que interviene PEI.",
    "ETAPAS": "Las partes en que se divide cada proyecto.",
    "FRENTES": "Los frentes de trabajo de cada etapa y qué contratista está a cargo, con fechas. Es lo que permite saber quién era responsable en la fecha de cada archivo.",
    "PERSONAS": "El equipo que sube archivos. Con esto el sistema sabe quién subió cada uno y quién debe revisarlo.",
}


def generar_plantilla(ruta: Path) -> Path:
    """Crea la tabla maestra vacía, con una fila de ejemplo por hoja, notas
    en cada encabezado y una hoja de instrucciones para el director."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    ejemplo = {
        "PROYECTOS": [["P001", "Alejandría (ejemplo, reemplazar)", "ACTIVO", "2025-05-01", "", "alejandria, calle 100"]],
        "ETAPAS": [["P001", "E04", "Etapa 4 - Sector 5", "2025-05-01", "", "K0+000", "K0+800",
                    "", "", "", "", "etapa 4, sector 5"]],
        "FRENTES": [["P001", "E04", "F01", "Redes sanitarias", "EMS", "2025-05-01", ""]],
        "CONTRATISTAS": [["EMS", "EMS (ejemplo)", "obra@ejemplo.com"]],
        "PERSONAS": [["topografo@ejemplo.com", "Nombre del topógrafo", "topografo", "PEI"]],
    }
    azul = PatternFill("solid", fgColor="1F4E78")
    gris = PatternFill("solid", fgColor="D9E1F2")
    blanco = Font(bold=True, color="FFFFFF")
    negrita = Font(bold=True)
    ajuste = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ins = wb.active
    ins.title = "INSTRUCCIONES"
    for col, ancho in zip("ABCDE", (16, 18, 70, 24, 13)):
        ins.column_dimensions[col].width = ancho

    def linea(texto="", estilo=None):
        ins.append([texto])
        c = ins.cell(row=ins.max_row, column=1)
        if estilo:
            c.font = estilo

    linea("TABLA MAESTRA PEI", Font(bold=True, size=14))
    linea("Es la lista oficial de proyectos, etapas, frentes, contratistas y personas de PEI. El sistema la "
          "lee cada vez que procesa archivos: con ella revisa que los nombres de los archivos sean correctos, "
          "crea las carpetas en Drive y sugiere cómo clasificar los archivos antiguos.")
    linea("La llena y la actualiza solo el director o su suplente. El resto del equipo solo puede verla.")
    linea("Se actualiza cuando empieza un proyecto, una etapa o un frente; cuando un frente cambia de "
          "contratista, y cuando una persona entra o sale del equipo.")
    linea()
    linea("ORDEN PARA LLENARLA", negrita)
    for i, (hoja, texto) in enumerate(PROPOSITO_HOJAS.items(), 1):
        linea(f"{i}. {hoja}. {texto}")
    linea()
    linea("QUÉ VA EN CADA COLUMNA (también aparece como nota al pasar el cursor por cada encabezado)", negrita)
    ins.append(["Hoja", "Columna", "Qué escribir", "Ejemplo", "Obligatoria"])
    for c in ins[ins.max_row]:
        c.font, c.fill = blanco, azul
    hoja_anterior = None
    for hoja, columna, texto, ej, obligatoria in GUIA_COLUMNAS:
        ins.append([hoja if hoja != hoja_anterior else "", columna, texto, ej, obligatoria])
        fila = ins[ins.max_row]
        for c in fila:
            c.alignment = ajuste
        if hoja != hoja_anterior:
            for c in fila:
                c.fill = gris
        hoja_anterior = hoja
    linea()
    linea("REGLAS", negrita)
    for regla in (
        "1. Los códigos (P001, E04, F01, EMS) no se cambian ni se reutilizan una vez creados, porque ya están en los nombres de los archivos.",
        "2. No se borran filas de FRENTES. Si un frente cambia de contratista, se pone fecha fin en la fila actual y se agrega una nueva.",
        "3. Las fechas van como AAAA-MM-DD (por ejemplo 2025-05-01). Fecha fin vacía significa que sigue vigente.",
        "4. Un archivo que mencione un proyecto, etapa o frente que no esté aquí será rechazado con el motivo.",
        "5. ZZ significa general o varios: se usa en el nombre del archivo cuando no aplica una etapa o un frente.",
        "6. PEI es el origen de los archivos que produce la interventoría; el código del contratista, de los que entrega él.",
        "7. Borre las filas de ejemplo antes de empezar.",
    ):
        linea(regla)
    linea()
    linea("TIPOS DE ARCHIVO VÁLIDOS EN EL NOMBRE", negrita)
    for codigo, desc in TIPOS.items():
        ins.append([codigo, desc])
    for fila in ins.iter_rows(min_row=1, max_row=8):
        fila[0].alignment = Alignment(wrap_text=False)

    notas = {}
    for hoja, columna, texto, ej, obligatoria in GUIA_COLUMNAS:
        notas[(hoja, columna)] = f"{texto} Ejemplo: {ej or '(vacío)'}. Obligatoria: {obligatoria}."
    for hoja, columnas in HOJAS.items():
        ws = wb.create_sheet(hoja)
        ws.append(columnas)
        for c in ws[1]:
            c.font, c.fill = blanco, azul
            nota = notas.get((hoja, c.value))
            if nota:
                c.comment = Comment(nota, "PEI", width=320, height=120)
        for fila in ejemplo[hoja]:
            ws.append(fila)
        for i, _ in enumerate(columnas, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
        ws.freeze_panes = "A2"
        if hoja == "PROYECTOS":
            dv = DataValidation(type="list", formula1='"ACTIVO,TERMINADO,SUSPENDIDO"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add("C2:C200")
        if hoja == "PERSONAS":
            dv = DataValidation(type="list", formula1='"director,suplente,residente,topografo,auxiliar"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add("C2:C200")
    # Orden de las pestañas igual al orden de llenado
    wb._sheets = [wb["INSTRUCCIONES"]] + [wb[h] for h in PROPOSITO_HOJAS]
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta
